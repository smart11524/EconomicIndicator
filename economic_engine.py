"""
Economic Engine for Petroleum Field Development Evaluation
Refactored from Jupyter Notebook model with end-to-end analysis, 
figure generation, Excel export, and ZIP packaging.
"""

import os
import re
import sys
import io
import shutil
import datetime
import zipfile
import tempfile
import warnings
from dataclasses import dataclass, field, replace, asdict
from typing import Dict, Any, List, Optional, Tuple

import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")  # Non-interactive backend for web/server use
import matplotlib.pyplot as plt
from openpyxl import load_workbook

try:
    import xlrd
except ImportError:
    xlrd = None

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

# Standard configuration defaults
DEFAULT_CONFIG = {
    "APPLY_PRICE_ESCALATION": False,
    "ECONOMIC_LIMIT_CUTOFF": False,
    "LOSS_CARRY_FORWARD": False,
    "DB_FACTOR": 1.5,
    "DDB_FACTOR": 2.0,
    "PAYBACK_LIMIT_YEARS": 10,
    "STARTUP_DELAYS": [0, 1, 2, 3, 4, 5],
    "INFLATE_DURING_DELAY": True,
    "CAPEX_ESCALATION": 0.0,
    "DEVELOPMENT_SCALES": [0.6, 0.8, 1.0, 1.2, 1.5, 2.0],
    "FACILITY_SCALING_EXPONENT": 0.6,
    "RESERVOIR_RESPONSE_EXPONENT": 0.85,
    "CUSTOM_ALTERNATIVES": {},
    "CAPITAL_BUDGET": None,
    "NPV_PROFILE_MAX_RATE": 0.50,
    "SAVE_RESULTS": True,
    "RESULTS_FILE": "Economic_Indicators_Results.xlsx",
}

@dataclass
class ProjectInputs:
    """All inputs of one field-development project (units follow the template)."""
    name: str = "Template case"
    # production
    q1_mbd: float = 0.0          # oil rate in year 1, thousand barrels/day (MBD)
    decline: float = 0.0         # annual decline (fraction)
    gor: float = 0.0             # gas-oil ratio, SCF/STB
    # investment
    wells: float = 0.0
    drill_cost: float = 0.0      # $MM/well
    pipe_len: float = 0.0        # km
    d_crude: float = 0.0         # inches
    d_gas: float = 0.0           # inches
    pipe_unit_cost: float = 0.0  # $/inch-km
    facilities: float = 0.0      # $MM
    camp: float = 0.0            # $MM
    # operating cost (year-0 basis) and inflation (fractions per year)
    fuel: float = 0.0;        fuel_infl: float = 0.0          # $MM/yr
    spares: float = 0.0;      spares_infl: float = 0.0        # $MM/yr
    maint: float = 0.0;       maint_infl: float = 0.0         # $MM/yr
    salary: float = 0.0;      salary_infl: float = 0.0        # $/person/month
    catering: float = 0.0;    catering_infl: float = 0.0      # $/person/day
    overhead: float = 0.0;    overhead_infl: float = 0.0      # $MM/yr
    well_serv: float = 0.0;   well_serv_infl: float = 0.0     # $MM/yr/well
    personnel: float = 0.0
    # cash-flow data
    oil_price: float = 0.0;   oil_esc: float = 0.0            # $/bbl , escalation fraction/yr
    gas_price: float = 0.0;   gas_esc: float = 0.0            # $/MCF
    tax: float = 0.0                                          # fraction
    disc: float = 0.0                                         # discount rate, fraction
    # depreciation
    dep_method: str = "SL"       # SL, DBwB, DB, DDBwB, DDB, SYD, UOP
    dep_years: int = 10
    salvage: float = 0.0         # $MM
    life: int = 20               # project life, years
    # extras read from the template
    sens: dict = field(default_factory=dict)              # {"Production Rate": (lo, hi), ...} fractions
    template_results: dict = field(default_factory=dict)  # values already computed inside the workbook


_INPUT_MAP = [
    ("q1_mbd", "Oil production rate in year 1", 4, 1), ("decline", "Decline rate", 4, 0.01),
    ("gor", "Gas Oil Ratio", 4, 1),
    ("wells", "Number of wells", 4, 1), ("drill_cost", "Drilling cost", 4, 1),
    ("pipe_len", "Length of pipelines", 4, 1), ("d_crude", "Diameter of crude pipeline", 4, 1),
    ("d_gas", "Diameter of gas pipeline", 4, 1), ("pipe_unit_cost", "Cost of pipeline", 4, 1),
    ("facilities", "Cost of production facilities", 4, 1), ("camp", "Cost of camp", 4, 1),
    ("fuel", "Fuel", 4, 1), ("spares", "Spare parts", 4, 1), ("maint", "Maintenance Requirements", 4, 1),
    ("salary", "Salaries", 4, 1), ("catering", "Catering & Services", 4, 1),
    ("overhead", "Overhead costs", 4, 1), ("well_serv", "Well services & workovers", 4, 1),
    ("personnel", "Number of personnel", 4, 1),
    ("oil_price", "Crude selling Price", 4, 1), ("gas_price", "Gas selling Price", 4, 1),
    ("tax", "Tax rate", 4, 0.01), ("disc", "Discount Rate", 4, 0.01),
    ("dep_years", "Number of Years", 4, 1),
]

_RATE_MAP = [
    ("fuel_infl", "Fuel"), ("spares_infl", "Spare parts"), ("maint_infl", "Maintenance Requirements"),
    ("salary_infl", "Salaries"), ("catering_infl", "Catering & Services"), ("overhead_infl", "Overhead costs"),
    ("well_serv_infl", "Well services & workovers"), ("oil_esc", "Crude selling Price"), ("gas_esc", "Gas selling Price"),
]

_DEP_HEADERS = {"str. line": "SL", "dbw/b": "DBwB", "db": "DB", "ddbw/b": "DDBwB",
                "ddb": "DDB", "s.y.d": "SYD", "unt. prod": "UOP"}

DEP_NAMES = {"SL": "Straight line", "DBwB": "Declining balance w/ balloon", "DB": "Declining balance",
             "DDBwB": "Double declining w/ balloon", "DDB": "Double declining balance",
             "SYD": "Sum-of-years-digits", "UOP": "Unit of production"}


def _norm(s):
    """Normalize label: lowercase, single spaces."""
    return re.sub(r"\s+", " ", str(s)).strip().lower()


class _XlsSheet:
    """Minimal wrapper so an .xls sheet (xlrd) behaves like openpyxl."""
    class _Cell:
        __slots__ = ("value",)
        def __init__(self, value):
            self.value = value
    def __init__(self, sheet):
        self._s = sheet
        self.max_row = sheet.nrows
    def cell(self, r, c):
        if r > self._s.nrows or c > self._s.ncols:
            return self._Cell(None)
        cell = self._s.cell(r - 1, c - 1)
        if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR) or cell.value == "":
            return self._Cell(None)
        return self._Cell(cell.value)


def open_input_sheet(path):
    """Open the 'Input&Output' sheet of an .xlsx/.xlsm (openpyxl) or raw .xls (xlrd)."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xls":
        if xlrd is None:
            raise ImportError("xlrd is required to read .xls files. Please install xlrd.")
        wb = xlrd.open_workbook(path)
        names = wb.sheet_names()
        get = lambda n: _XlsSheet(wb.sheet_by_name(n))
    else:
        wb = load_workbook(path, data_only=True)
        names = wb.sheetnames
        get = lambda n: wb[n]
    sheet = next((n for n in names if _norm(n) == "input&output"), None)
    if sheet is None:
        raise ValueError(f"Sheet 'Input&Output' not found. Available sheets: {names}. "
                         "Please ensure the uploaded workbook adheres to the Economic Evaluation Study template.")
    return get(sheet)


def read_template(path) -> ProjectInputs:
    """Read the template and return a populated ProjectInputs object."""
    ws = open_input_sheet(path)

    rows = {}
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v is not None:
            rows.setdefault(_norm(v), []).append(r)

    def row_of(label, after=0):
        for r in rows.get(_norm(label), []):
            if r > after:
                return r
        return None

    def number(cell):
        return float(cell) if isinstance(cell, (int, float)) and not isinstance(cell, bool) else None

    p, missing = ProjectInputs(), []
    for attr, label, col, mult in _INPUT_MAP:
        r = row_of(label)
        val = number(ws.cell(r, col).value) if r else None
        if val is None:
            missing.append(f"{label!r} (expected in row {r}, column {col})")
        else:
            setattr(p, attr, val * mult)

    for attr, label in _RATE_MAP:
        r = row_of(label)
        val = number(ws.cell(r, 7).value) if r else None
        setattr(p, attr, (val or 0.0) / 100.0)

    if missing:
        raise ValueError("Dataset does not match the template. Problem inputs:\n  - " + "\n  - ".join(missing))

    r = row_of("Salvage Value")
    p.salvage = (number(ws.cell(r, 4).value) or 0.0) if r else 0.0

    r = row_of("Depreciation Method")
    chosen = []
    if r:
        for c in range(4, 11):
            hdr = _norm(ws.cell(r, c).value or "")
            mark = ws.cell(r + 1, c).value
            if hdr in _DEP_HEADERS and mark == 1:
                chosen.append(_DEP_HEADERS[hdr])
    if len(chosen) != 1:
        chosen = ["SL"]
    p.dep_method = chosen[0]
    p.dep_years = int(p.dep_years)

    p.life = 20
    for r_ in range(1, ws.max_row + 1):
        m = re.search(r"project\s*life\s*=\s*(\d+)", str(ws.cell(r_, 1).value or ""), re.I)
        if m:
            p.life = int(m.group(1))
            break

    hdr_row = next((r_ for r_ in range(1, ws.max_row + 1)
                    if re.match(r"\s*s[ie]n[sc]?[is]tiv", str(ws.cell(r_, 1).value or ""), re.I)), None)
    if hdr_row:
        for label in ("Production Rate", "Capex", "Opex", "Oil Price"):
            r = row_of(label, after=hdr_row)
            lo, hi = (number(ws.cell(r, 3).value), number(ws.cell(r, 5).value)) if r else (None, None)
            if lo is not None and hi is not None:
                p.sens[label] = (lo / 100.0, hi / 100.0)

    for key, label in (("NPV", "Net Present Value (NPV)"), ("IRR", "Internal Rate of Return (IRR)"),
                       ("Payback", "Payback Period (Payout)"), ("CPI", "Capital Productivity Index (CPI)")):
        for r_ in range(1, ws.max_row + 1):
            if _norm(ws.cell(r_, 4).value or "") == _norm(label):
                val = number(ws.cell(r_, 8).value)
                if val is not None:
                    p.template_results[key] = val
    if len(p.template_results) < 4:
        p.template_results = {}

    m = re.search(r"\(([^)]*)\)", str(ws.cell(2, 1).value or ""))
    p.name = (m.group(1).strip() if m and m.group(1).strip() else "Base")
    if p.life < 1 or p.dep_years < 1:
        raise ValueError("Project life and depreciation years must be >= 1.")
    return p


def capex_breakdown(p: ProjectInputs) -> Dict[str, float]:
    """Investment ($MM): drilling + pipelines + facilities + camp (all spent at year 0)."""
    drilling = p.wells * p.drill_cost
    pipelines = p.pipe_len * (p.d_crude + p.d_gas) * p.pipe_unit_cost / 1e6
    return {
        "Drilling": float(drilling),
        "Pipelines": float(pipelines),
        "Production facilities": float(p.facilities),
        "Camp": float(p.camp),
        "TOTAL CAPEX": float(drilling + pipelines + p.facilities + p.camp)
    }


def depreciation_schedule(capex, method, n, salvage, production=None, life=None,
                          db_factor=1.5, ddb_factor=2.0) -> np.ndarray:
    """Depreciation charge for project years 0..life (index 0 = 0)."""
    life = life if life is not None else n
    n = int(min(n, life))
    dep = np.zeros(life + 1)
    base = max(capex - salvage, 0.0)
    if method == "SL":
        dep[1:n + 1] = base / n
    elif method in ("DB", "DBwB", "DDB", "DDBwB"):
        rate = (db_factor if method.startswith("DB") else ddb_factor) / n
        book = capex
        for t in range(1, n + 1):
            d = min(rate * book, book - salvage)
            if method.endswith("wB") and t == n:
                d = book - salvage
            dep[t] = max(d, 0.0)
            book -= dep[t]
    elif method == "SYD":
        syd = n * (n + 1) / 2.0
        for t in range(1, n + 1):
            dep[t] = base * (n - t + 1) / syd
    elif method == "UOP":
        prod = np.asarray(production, float)
        prod[0] = 0.0
        dep[1:] = base * prod[1:] / prod[1:].sum() if prod[1:].sum() > 0 else 0.0
    else:
        raise ValueError(f"Unknown depreciation method {method!r}")
    return dep


def build_cash_flow(p: ProjectInputs, startup_delay: int = 0, mult: Optional[dict] = None, cfg: Optional[dict] = None) -> pd.DataFrame:
    """Annual cash-flow table ($MM) for project inputs `p`."""
    cfg = cfg or DEFAULT_CONFIG
    m = {"prod": 1.0, "capex": 1.0, "opex": 1.0, "price": 1.0, **(mult or {})}
    d, N = int(startup_delay), int(p.life)
    year = np.arange(0, N + d + 1)
    tp = year - d
    producing = tp >= 1
    idx = np.maximum(year if cfg.get("INFLATE_DURING_DELAY", True) else tp, 0)
    pidx = idx if cfg.get("APPLY_PRICE_ESCALATION", False) else np.zeros_like(idx)

    # Production
    oil = np.where(producing, p.q1_mbd * 365 / 1000 * m["prod"] * (1 - p.decline) ** (np.maximum(tp, 1) - 1), 0.0)
    gas = oil * p.gor / 1000.0

    # Revenue
    oil_rev = p.oil_price * m["price"] * (1 + p.oil_esc) ** pidx * oil
    gas_rev = p.gas_price * (1 + p.gas_esc) ** pidx * gas
    revenue = oil_rev + gas_rev

    # Operating costs
    g = lambda base, infl: np.where(producing, base * (1 + infl) ** idx, 0.0) * m["opex"]
    items = {
        "Fuel": g(p.fuel, p.fuel_infl),
        "Spares": g(p.spares, p.spares_infl),
        "Maintenance": g(p.maint, p.maint_infl),
        "Salaries": g(p.salary * p.personnel * 12 / 1e6, p.salary_infl),
        "Catering": g(p.catering * p.personnel * 365 / 1e6, p.catering_infl),
        "Overhead": g(p.overhead, p.overhead_infl),
        "Well services": g(p.wells * p.well_serv, p.well_serv_infl),
    }
    opex = sum(items.values())

    if cfg.get("ECONOMIC_LIMIT_CUTOFF", False):
        neg = np.where(producing & (revenue - opex < 0))[0]
        if len(neg):
            k = neg[0]
            for arr in (oil, gas, oil_rev, gas_rev, revenue, opex, *items.values()):
                arr[k:] = 0.0

    capex_total = capex_breakdown(p)["TOTAL CAPEX"] * m["capex"] * (1 + cfg.get("CAPEX_ESCALATION", 0.0)) ** d
    capex = np.where(tp == 0, capex_total, 0.0)

    dep = np.zeros(len(year))
    dep[d:] = depreciation_schedule(capex_total, p.dep_method, p.dep_years, p.salvage,
                                    production=oil[d:], life=N,
                                    db_factor=cfg.get("DB_FACTOR", 1.5), ddb_factor=cfg.get("DDB_FACTOR", 2.0))

    taxable = revenue - opex - dep
    if cfg.get("LOSS_CARRY_FORWARD", False):
        tax, loss = np.zeros(len(year)), 0.0
        for i, ti in enumerate(taxable):
            if ti < 0:
                loss += -ti
            else:
                used = min(loss, ti)
                loss -= used
                tax[i] = (ti - used) * p.tax
    else:
        tax = np.maximum(0.0, taxable) * p.tax

    ncf = revenue - capex - opex - tax
    out = pd.DataFrame({
        "Year": year, "ProjYear": tp, "Oil_MMB": oil, "Gas_BCF": gas,
        "Oil_Rev": oil_rev, "Gas_Rev": gas_rev, "Revenue": revenue, "CAPEX": capex,
        **items, "OPEX": opex, "Depreciation": dep, "Taxable_Income": taxable,
        "Tax": tax, "NCF": ncf
    })
    out["Cum_NCF"] = out["NCF"].cumsum()
    return out


def discount_factor(rate: float, t) -> np.ndarray:
    return 1.0 / (1.0 + rate) ** np.asarray(t, dtype=float)


def present_value(cash_flow, rate: float, t) -> np.ndarray:
    return np.asarray(cash_flow, dtype=float) * discount_factor(rate, t)


def npv(cash_flow, rate: float, t=None) -> float:
    cash_flow = np.asarray(cash_flow, dtype=float)
    t = np.arange(len(cash_flow)) if t is None else t
    return float(present_value(cash_flow, rate, t).sum())


def sign_changes(cash_flow) -> int:
    s = np.sign(np.asarray(cash_flow, float))
    s = s[s != 0]
    return int((s[1:] != s[:-1]).sum())


def irr_roots(cash_flow) -> List[float]:
    cf = np.asarray(cash_flow, dtype=float)
    nz = np.nonzero(cf)[0]
    if len(nz) < 2:
        return []
    cf = cf[nz[0]:nz[-1] + 1]
    if not (cf.min() < 0 < cf.max()):
        return []
    t = np.arange(len(cf))
    roots = []
    for z in np.roots(cf[::-1]):
        if abs(z.imag) > 1e-6 * max(1.0, abs(z)) or z.real <= 0:
            continue
        r = 1.0 / z.real - 1.0
        if r <= -0.99:
            continue
        for _ in range(50):
            f = (cf / (1 + r) ** t).sum()
            fp = (-t * cf / (1 + r) ** (t + 1)).sum()
            if fp == 0:
                break
            step = f / fp
            r -= step
            if abs(step) < 1e-13:
                break
        if abs((cf / (1 + r) ** t).sum()) < 1e-6 * max(1.0, np.abs(cf).max()):
            if all(abs(r - x) > 1e-7 for x in roots):
                roots.append(float(r))
    return sorted(roots)


def mirr(cash_flow, rate: float) -> float:
    cf = np.asarray(cash_flow, dtype=float)
    nz = np.nonzero(cf)[0]
    if len(nz) < 2:
        return float("nan")
    cf = cf[nz[0]:nz[-1] + 1]
    n, t = len(cf) - 1, np.arange(len(cf))
    fv_pos = (np.where(cf > 0, cf, 0.0) * (1 + rate) ** (n - t)).sum()
    pv_neg = -(np.where(cf < 0, cf, 0.0) / (1 + rate) ** t).sum()
    return float((fv_pos / pv_neg) ** (1.0 / n) - 1.0) if fv_pos > 0 and pv_neg > 0 else float("nan")


def payback_period(cash_flow, times=None) -> Optional[float]:
    cf = np.asarray(cash_flow, dtype=float)
    times = np.arange(len(cf)) if times is None else np.asarray(times, dtype=float)
    cum = np.cumsum(cf)
    for i in range(len(cf) - 1):
        if cum[i] < 0 <= cum[i + 1]:
            return float(times[i] + (-cum[i]) / cf[i + 1] * (times[i + 1] - times[i]))
    return None


def discounted_payback_period(cash_flow, rate: float, times=None) -> Optional[float]:
    cf = np.asarray(cash_flow, dtype=float)
    times = np.arange(len(cf)) if times is None else np.asarray(times, dtype=float)
    return payback_period(present_value(cf, rate, times), times)


def profitability(cf: pd.DataFrame, rate: float) -> Dict[str, float]:
    dfs = discount_factor(rate, cf["Year"])
    pv_capex = float((cf["CAPEX"] * dfs).sum())
    pv_net_inflow = float(((cf["Revenue"] - cf["OPEX"] - cf["Tax"]) * dfs).sum())
    npv_ = pv_net_inflow - pv_capex
    return {
        "PV_Capex": pv_capex,
        "PV_NetInflow": pv_net_inflow,
        "PI": pv_net_inflow / pv_capex if pv_capex > 0 else float("nan"),
        "PIR": npv_ / pv_capex if pv_capex > 0 else float("nan"),
        "PIR_undiscounted": float(cf["NCF"].sum() / cf["CAPEX"].sum()) if cf["CAPEX"].sum() > 0 else float("nan")
    }


def fmt_roots(roots: List[float]) -> str:
    return " | ".join(f"{r:.2%}" for r in roots) if roots else "none"


def evaluate(cf: pd.DataFrame, rate: float) -> Dict[str, Any]:
    proj = cf["ProjYear"].values
    on = proj >= 0
    ncf, yr = cf["NCF"].values, cf["Year"].values
    ncf_p, t_p = ncf[on], proj[on].astype(float)
    m = {
        "Capex": float(cf["CAPEX"].sum()),
        "NPV": npv(ncf, rate, yr),
        "IRR_roots": irr_roots(ncf_p),
        "MIRR": mirr(ncf_p, rate),
        "Payback": payback_period(ncf_p, t_p),
        "Disc_Payback": discounted_payback_period(ncf_p, rate, t_p),
        "Sign_Changes": sign_changes(ncf_p)
    }
    m.update(profitability(cf, rate))
    m["N_IRR"] = len(m["IRR_roots"])
    m["IRR_text"] = fmt_roots(m["IRR_roots"])
    m["IRR_low"] = m["IRR_roots"][0] if m["IRR_roots"] else float("nan")
    m["IRR_high"] = m["IRR_roots"][-1] if m["IRR_roots"] else float("nan")
    m["Last_Producing_Year"] = int(cf.loc[cf["Revenue"] > 0, "ProjYear"].max()) if (cf["Revenue"] > 0).any() else 0
    return m


def screening_table(m: Dict[str, Any], rate: float, cfg: Optional[dict] = None) -> pd.DataFrame:
    cfg = cfg or DEFAULT_CONFIG
    limit_yr = cfg.get("PAYBACK_LIMIT_YEARS", 10)
    if m["N_IRR"] == 1:
        irr_val, irr_ok = f"{m['IRR_low']:.2%}", m["IRR_low"] > rate
        irr_crit = f"IRR > {rate:.0%}"
    elif m["N_IRR"] == 0:
        irr_val, irr_ok, irr_crit = "no IRR", False, "cash flow has no sign change"
    else:
        irr_val = m["IRR_text"]
        irr_ok = (m["NPV"] > 0) and (m["MIRR"] > rate)
        irr_crit = f"multiple IRRs -> NPV > 0 and MIRR ({m['MIRR']:.2%}) > {rate:.0%}"
    
    pb_ok = m["Payback"] is not None and m["Payback"] < limit_yr
    dpb_ok = m["Disc_Payback"] is not None
    rows = [
        ("NPV ($MM)", f"{m['NPV']:,.3f}", "NPV > 0", m["NPV"] > 0),
        ("IRR", irr_val, irr_crit, irr_ok),
        ("Payback (years)", "never" if m["Payback"] is None else f"{m['Payback']:.2f}",
         f"< {limit_yr} years", pb_ok),
        ("Discounted payback (years)", "never" if m["Disc_Payback"] is None else f"{m['Disc_Payback']:.2f}",
         "recovered within project life", dpb_ok),
        ("Profitability index PI", f"{m['PI']:.3f}", "PI > 1", m["PI"] > 1),
        ("Profit/investment ratio PIR (= CPI)", f"{m['PIR']:.3f}", "PIR > 0", m["PIR"] > 0)
    ]
    df = pd.DataFrame(rows, columns=["Indicator", "Value", "Criterion", "OK?"])
    df["OK?"] = np.where(df["OK?"], "OK", "Not OK")
    return df


def make_alternative(base: ProjectInputs, scale: float, label: Optional[str] = None, cfg: Optional[dict] = None) -> ProjectInputs:
    cfg = cfg or DEFAULT_CONFIG
    fac_exp = cfg.get("FACILITY_SCALING_EXPONENT", 0.6)
    res_exp = cfg.get("RESERVOIR_RESPONSE_EXPONENT", 0.85)
    wells = max(1, round(base.wells * scale))
    sw = wells / base.wells
    sf = sw ** fac_exp
    sq = sw ** res_exp
    alt = replace(base, name=label or f"{scale:.1f}x wells", wells=wells, q1_mbd=base.q1_mbd * sq,
                  facilities=base.facilities * sf, camp=base.camp * sf, pipe_unit_cost=base.pipe_unit_cost * sf,
                  fuel=base.fuel * sf, spares=base.spares * sf, maint=base.maint * sf,
                  overhead=base.overhead * sf, personnel=max(1, round(base.personnel * sf)))
    return alt


def build_alternatives(base: ProjectInputs, cfg: Optional[dict] = None) -> Dict[str, ProjectInputs]:
    cfg = cfg or DEFAULT_CONFIG
    alts = {}
    custom = cfg.get("CUSTOM_ALTERNATIVES", {})
    if custom:
        for label, overrides in custom.items():
            alts[label] = replace(base, name=label, **overrides)
        alts.setdefault(f"{base.name} (template)", base)
    else:
        scales = cfg.get("DEVELOPMENT_SCALES", [0.6, 0.8, 1.0, 1.2, 1.5, 2.0])
        for s in scales:
            label = f"{base.name} (template)" if abs(s - 1.0) < 1e-9 else f"{s:.1f}x wells ({round(base.wells * s)} wells)"
            alts[label] = base if abs(s - 1.0) < 1e-9 else make_alternative(base, s, label, cfg)
    return alts


def incremental_chain(cf_dict: Dict[str, pd.DataFrame], rate: float, order: List[str]) -> Tuple[pd.DataFrame, str]:
    rows, current = [], order[0]
    for nxt in order[1:]:
        a, b = cf_dict[current], cf_dict[nxt]
        n = max(len(a), len(b))
        d_ncf = np.zeros(n)
        d_ncf[:len(b)] += b["NCF"].values
        d_ncf[:len(a)] -= a["NCF"].values
        d_year = np.arange(n)
        d_capex = float(b["CAPEX"].sum() - a["CAPEX"].sum())
        d_npv = npv(d_ncf, rate, d_year)
        d_roots = irr_roots(d_ncf)
        d_mirr = mirr(d_ncf, rate)
        accept = d_npv > 0
        rows.append({
            "Increment": f"{current}  ->  {nxt}",
            "Delta Capex ($MM)": d_capex,
            "Delta NPV ($MM)": d_npv,
            "Delta IRR": fmt_roots(d_roots),
            "Delta MIRR": d_mirr,
            "Delta PIR": d_npv / d_capex if d_capex != 0 else float("nan"),
            "Decision": "TAKE increment (move up)" if accept else "STOP (stay with smaller)"
        })
        if accept:
            current = nxt
    return pd.DataFrame(rows), current


def _slug(text, default="item"):
    s = re.sub(r"[^A-Za-z0-9]+", "_", str(text)).strip("_")
    return (s[:60] or default)


def run_full_analysis(file_path: str, output_base_dir: Optional[str] = None, user_config: Optional[dict] = None) -> Dict[str, Any]:
    """
    Executes the comprehensive economic evaluation on the uploaded Excel file.
    Creates tables, high-resolution figures, results workbooks, and a complete ZIP archive.
    """
    cfg = {**DEFAULT_CONFIG, **(user_config or {})}
    
    # Generate run ID and output directory
    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    stem = os.path.splitext(os.path.basename(file_path))[0]
    run_id = f"{_slug(stem)}_{timestamp}"
    
    if output_base_dir is None:
        output_dir = os.path.join(tempfile.gettempdir(), f"econ_run_{run_id}")
    else:
        output_dir = os.path.join(output_base_dir, f"run_{run_id}")
        
    fig_dir = os.path.join(output_dir, "figures")
    table_dir = os.path.join(output_dir, "tables")
    os.makedirs(fig_dir, exist_ok=True)
    os.makedirs(table_dir, exist_ok=True)
    
    log_stream = io.StringIO()
    def log(msg):
        print(msg)
        log_stream.write(msg + "\n")
        
    log(f"Economic Evaluation Started: {datetime.datetime.now():%Y-%m-%d %H:%M:%S}")
    log(f"File: {file_path}")
    
    # 1. Parse Excel Template
    base_inputs = read_template(file_path)
    disc = cfg.get("DISCOUNT_RATE_OVERRIDE") or base_inputs.disc
    
    log(f"Project Name: {base_inputs.name}")
    log(f"Project Life: {base_inputs.life} years | Discount Rate: {disc:.1%} | Tax: {base_inputs.tax:.0%}")
    log(f"Initial Oil Rate: {base_inputs.q1_mbd} MBD | Decline: {base_inputs.decline:.1%} | GOR: {base_inputs.gor} SCF/STB")
    log(f"Depreciation: {DEP_NAMES[base_inputs.dep_method]} over {base_inputs.dep_years} years")
    
    capex_dict = capex_breakdown(base_inputs)
    capex_df = pd.Series(capex_dict, name="$MM").to_frame()
    capex_df.to_csv(os.path.join(table_dir, "capex_breakdown.csv"))
    
    # 2. Build Base Cash Flow
    base_cf = build_cash_flow(base_inputs, cfg=cfg)
    base_cf.round(4).to_csv(os.path.join(table_dir, "base_cash_flow.csv"), index=False)
    
    base_metrics = evaluate(base_cf, disc)
    screening_df = screening_table(base_metrics, disc, cfg)
    screening_df.to_csv(os.path.join(table_dir, "screening_table.csv"), index=False)
    
    # 3. Workbook Verification Check
    ver_df = None
    if base_inputs.template_results:
        tr = base_inputs.template_results
        closest_root = min(base_metrics["IRR_roots"], key=lambda r: abs(r - tr["IRR"])) if base_metrics["IRR_roots"] and "IRR" in tr else float("nan")
        checks = [
            ("NPV ($MM)", base_metrics["NPV"], tr.get("NPV")),
            ("IRR (matching root)", closest_root, tr.get("IRR")),
            ("Payback (years)", base_metrics["Payback"], tr.get("Payback")),
            ("CPI = NPV/CAPEX", base_metrics["Capex"] and base_metrics["NPV"] / base_metrics["Capex"], tr.get("CPI"))
        ]
        ver_df = pd.DataFrame([
            (n, py, xl, (py - xl) if (py is not None and xl is not None) else np.nan)
            for n, py, xl in checks
        ], columns=["Indicator", "Python", "Workbook", "Difference"])
        ver_df["Match"] = np.where(ver_df["Difference"].abs() < 1e-4, "YES", "NO")
        ver_df.to_csv(os.path.join(table_dir, "workbook_verification.csv"), index=False)
        log("Verification against workbook completed.")
    
    # 4. Economic Limit Cutoff Comparison
    cfg_fixed = {**cfg, "ECONOMIC_LIMIT_CUTOFF": False}
    cfg_limit = {**cfg, "ECONOMIC_LIMIT_CUTOFF": True}
    cf_fixed = build_cash_flow(base_inputs, cfg=cfg_fixed)
    cf_limit = build_cash_flow(base_inputs, cfg=cfg_limit)
    m_fixed = evaluate(cf_fixed, disc)
    m_limit = evaluate(cf_limit, disc)
    life_df = pd.DataFrame({
        "Fixed life (template)": {
            "Last producing year": m_fixed["Last_Producing_Year"], "NPV ($MM)": m_fixed["NPV"],
            "IRR root(s)": m_fixed["IRR_text"], "MIRR": f"{m_fixed['MIRR']:.2%}",
            "Payback (yr)": m_fixed["Payback"], "PIR": m_fixed["PIR"],
            "Undiscounted profit ($MM)": cf_fixed["NCF"].sum()
        },
        "Abandon at economic limit": {
            "Last producing year": m_limit["Last_Producing_Year"], "NPV ($MM)": m_limit["NPV"],
            "IRR root(s)": m_limit["IRR_text"], "MIRR": f"{m_limit['MIRR']:.2%}",
            "Payback (yr)": m_limit["Payback"], "PIR": m_limit["PIR"],
            "Undiscounted profit ($MM)": cf_limit["NCF"].sum()
        }
    }).T
    life_df.to_csv(os.path.join(table_dir, "economic_limit_comparison.csv"))
    
    # 5. Startup Delay Analysis
    startup_delays = cfg.get("STARTUP_DELAYS", [0, 1, 2, 3, 4, 5])
    startup_rows = []
    startup_cfs = {}
    for delay in startup_delays:
        cf_d = build_cash_flow(base_inputs, startup_delay=delay, cfg=cfg)
        m_d = evaluate(cf_d, disc)
        startup_cfs[delay] = cf_d
        startup_rows.append({
            "Delay (years)": delay, "NPV ($MM)": m_d["NPV"], "IRR": m_d["IRR_text"],
            "MIRR": m_d["MIRR"],
            "Payback from today (yr)": (m_d["Payback"] + delay) if m_d["Payback"] is not None else None,
            "Payback from start-up (yr)": m_d["Payback"], "PIR": m_d["PIR"]
        })
    startup_df = pd.DataFrame(startup_rows).set_index("Delay (years)")
    startup_df.round(4).to_csv(os.path.join(table_dir, "startup_delay.csv"))
    
    # 6. Development Alternatives & Incremental Analysis
    alternatives = build_alternatives(base_inputs, cfg=cfg)
    alt_cf = {label: build_cash_flow(p, cfg=cfg) for label, p in alternatives.items()}
    alt_metrics = {label: evaluate(cf_, disc) for label, cf_ in alt_cf.items()}
    
    summary_rows = []
    for label, m in alt_metrics.items():
        summary_rows.append({
            "Alternative": label, "Wells": alternatives[label].wells, "Capex ($MM)": m["Capex"],
            "NPV ($MM)": m["NPV"], "IRR": m["IRR_text"], "MIRR": m["MIRR"],
            "Payback (yr)": m["Payback"], "PI": m["PI"], "PIR": m["PIR"]
        })
    summary_df = pd.DataFrame(summary_rows).set_index("Alternative")
    cap_budget = cfg.get("CAPITAL_BUDGET")
    if cap_budget is not None:
        within = summary_df["Capex ($MM)"] <= cap_budget
        summary_df.insert(0, "Within budget", np.where(within, "YES", "NO"))
    summary_df.round(4).to_csv(os.path.join(table_dir, "development_alternatives.csv"))
    
    order_by_capex = summary_df["Capex ($MM)"].sort_values().index.tolist()
    if cap_budget is not None:
        order_by_capex = [a for a in order_by_capex if summary_df.loc[a, "Capex ($MM)"] <= cap_budget]
    inc_table, inc_winner = incremental_chain(alt_cf, disc, order_by_capex)
    inc_table.round(4).to_csv(os.path.join(table_dir, "incremental_analysis.csv"), index=False)
    
    # 7. Final Ranking & Recommendation
    rank_df = pd.DataFrame({
        "NPV rank": summary_df["NPV ($MM)"].rank(ascending=False),
        "MIRR rank": summary_df["MIRR"].rank(ascending=False),
        "Payback rank": summary_df["Payback (yr)"].rank(ascending=True),
        "PIR rank": summary_df["PIR"].rank(ascending=False),
    })
    rank_df["Average rank"] = rank_df.mean(axis=1)
    comparison = pd.concat([
        summary_df[["Capex ($MM)", "NPV ($MM)", "IRR", "MIRR", "Payback (yr)", "PIR"]],
        rank_df
    ], axis=1).sort_values("Average rank")
    comparison.round(4).to_csv(os.path.join(table_dir, "final_comparison.csv"))
    
    best_npv = summary_df["NPV ($MM)"].idxmax()
    recommendation = inc_winner if cap_budget is not None else best_npv
    
    log(f"Best by stand-alone NPV: {best_npv} (${summary_df.loc[best_npv, 'NPV ($MM)']:.2f} MM)")
    log(f"Incremental analysis winner: {inc_winner}")
    log(f"Final Recommendation: {recommendation}")
    
    # 8. Sensitivity Analysis
    mult_map = {"Production Rate": "prod", "Capex": "capex", "Opex": "opex", "Oil Price": "price"}
    best_p = alternatives[recommendation]
    sens_rows, spider = [], {}
    for label, (lo, hi) in base_inputs.sens.items():
        key = mult_map.get(label)
        if key is None:
            continue
        npv_lo = evaluate(build_cash_flow(best_p, mult={key: 1 + lo}, cfg=cfg), disc)["NPV"]
        npv_hi = evaluate(build_cash_flow(best_p, mult={key: 1 + hi}, cfg=cfg), disc)["NPV"]
        sens_rows.append({
            "Variable": label, "Low case": f"{lo:+.0%}", "NPV @ low": npv_lo,
            "High case": f"{hi:+.0%}", "NPV @ high": npv_hi, "Range ($MM)": abs(npv_hi - npv_lo)
        })
        grid = np.linspace(min(lo, 0), max(hi, 0), 11)
        spider[label] = (grid, [evaluate(build_cash_flow(best_p, mult={key: 1 + g}, cfg=cfg), disc)["NPV"] for g in grid])
        
    sens_df = pd.DataFrame(sens_rows).set_index("Variable").sort_values("Range ($MM)", ascending=False)
    sens_df.round(4).to_csv(os.path.join(table_dir, "sensitivity_analysis.csv"))
    
    # 9. Generate High-Res Figures
    generated_figures = {}
    
    # Fig 1: Discount Factors & PV Cash Flow
    fig, ax = plt.subplots(1, 2, figsize=(13, 4.5))
    rates_demo = [0.05, 0.10, 0.15, 0.20]
    for r in rates_demo:
        ax[0].plot(base_cf["Year"], discount_factor(r, base_cf["Year"]), label=f"{r:.0%}")
    ax[0].set(title="Discount Factor vs Year", xlabel="Year", ylabel="DF")
    ax[0].legend(title="Rate")
    w = 0.4
    dfs = discount_factor(disc, base_cf["Year"])
    pv_vals = base_cf["NCF"] * dfs
    ax[1].bar(base_cf["Year"] - w/2, base_cf["NCF"], w, label="Net Cash Flow (undiscounted)", color="#2b6cb0")
    ax[1].bar(base_cf["Year"] + w/2, pv_vals, w, label=f"Present Value @ {disc:.0%}", color="#dd6b20")
    ax[1].axhline(0, color="black", lw=0.8)
    ax[1].set(title="Undiscounted vs Present-Value Cash Flow", xlabel="Year", ylabel="$MM")
    ax[1].legend()
    plt.tight_layout()
    f1_path = os.path.join(fig_dir, "fig_01_discount_and_pv_flows.png")
    fig.savefig(f1_path, dpi=160, bbox_inches="tight")
    plt.close(fig)
    generated_figures["discount_and_pv"] = "figures/fig_01_discount_and_pv_flows.png"
    
    # Fig 2: NPV Component Build-up
    components = pd.Series({
        "PV of Revenue (+)": (base_cf["Revenue"] * dfs).sum(),
        "PV of CAPEX (-)": -(base_cf["CAPEX"] * dfs).sum(),
        "PV of OPEX (-)": -(base_cf["OPEX"] * dfs).sum(),
        "PV of Tax (-)": -(base_cf["Tax"] * dfs).sum(),
        "Net NPV": base_metrics["NPV"]
    })
    fig, ax = plt.subplots(figsize=(8, 4.5))
    colors = ["#2f855a", "#c53030", "#c53030", "#c53030", "#2b6cb0"]
    ax.bar(components.index, components.values, color=colors)
    ax.axhline(0, color="black", lw=0.8)
    ax.set_ylabel("$MM (Present Value)")
    ax.set_title(f"Build-up of NPV @ {disc:.0%}: ${base_metrics['NPV']:,.2f} MM")
    plt.xticks(rotation=20, ha="right")
    plt.tight_layout()
    f2_path = os.path.join(fig_dir, "fig_02_npv_component_buildup.png")
    fig.savefig(f2_path, dpi=160, bbox_inches="tight")
    plt.close(fig)
    generated_figures["npv_components"] = "figures/fig_02_npv_component_buildup.png"
    
    # Fig 3: NPV Profile & IRR Roots
    rate_grid = np.linspace(0.0, cfg.get("NPV_PROFILE_MAX_RATE", 0.50), 201)
    profile = [npv(base_cf["NCF"].values, r, base_cf["Year"]) for r in rate_grid]
    fig, ax = plt.subplots(figsize=(9, 4.8))
    ax.plot(rate_grid * 100, profile, color="#1a365d", lw=2.2, label="NPV Profile Curve")
    ax.axhline(0, color="black", lw=0.8)
    ax.axvline(disc * 100, color="#718096", ls="--", label=f"Project Rate {disc:.0%}")
    for k, r in enumerate(base_metrics["IRR_roots"]):
        ax.plot(r * 100, 0, "ro", ms=8, label="IRR Root(s)" if k == 0 else None)
        ax.annotate(f"{r:.2%}", (r * 100, 0), textcoords="offset points", xytext=(6, 10), fontweight="bold")
    ax.set(title=f"NPV Profile Curve (NPV vs Discount Rate)", xlabel="Discount Rate (%)", ylabel="NPV ($MM)")
    ax.legend()
    plt.tight_layout()
    f3_path = os.path.join(fig_dir, "fig_03_npv_profile.png")
    fig.savefig(f3_path, dpi=160, bbox_inches="tight")
    plt.close(fig)
    generated_figures["npv_profile"] = "figures/fig_03_npv_profile.png"
    
    # Fig 4: Cumulative Cash Flow & Payback
    cum_nc = np.cumsum(base_cf["NCF"].values)
    cum_pv = np.cumsum(pv_vals)
    fig, ax = plt.subplots(figsize=(9, 4.8))
    ax.plot(base_cf["Year"], cum_nc, "o-", color="#2b6cb0", label="Cumulative NCF (Undiscounted)")
    ax.plot(base_cf["Year"], cum_pv, "s--", color="#dd6b20", label=f"Cumulative Discounted NCF @ {disc:.0%}")
    ax.axhline(0, color="black", lw=0.8)
    if base_metrics["Payback"] is not None:
        ax.axvline(base_metrics["Payback"], color="#2b6cb0", ls=":")
        ax.annotate(f"Payback\n{base_metrics['Payback']:.2f} yr", (base_metrics["Payback"], cum_nc.min()*0.4),
                    textcoords="offset points", xytext=(6, 0), color="#2b6cb0", fontweight="bold")
    if base_metrics["Disc_Payback"] is not None:
        ax.axvline(base_metrics["Disc_Payback"], color="#dd6b20", ls=":")
        ax.annotate(f"Disc. Payback\n{base_metrics['Disc_Payback']:.2f} yr", (base_metrics["Disc_Payback"], cum_nc.min()*0.65),
                    textcoords="offset points", xytext=(6, 0), color="#dd6b20", fontweight="bold")
    ax.set(title="Payback Period - Cumulative Cash Flows", xlabel="Year", ylabel="$MM")
    ax.legend()
    plt.tight_layout()
    f4_path = os.path.join(fig_dir, "fig_04_payback_curves.png")
    fig.savefig(f4_path, dpi=160, bbox_inches="tight")
    plt.close(fig)
    generated_figures["payback_curves"] = "figures/fig_04_payback_curves.png"
    
    # Fig 5: Startup Delay Impact
    fig, ax = plt.subplots(1, 2, figsize=(13, 4.5))
    ax[0].plot(startup_df.index, startup_df["NPV ($MM)"], "o-", color="#1a365d", lw=2)
    ax[0].axhline(0, color="black", lw=0.7)
    ax[0].set(title="NPV vs Startup Delay", xlabel="Delay (years)", ylabel="NPV ($MM)")
    ax[1].plot(startup_df.index, startup_df["Payback from today (yr)"], "s-", color="#2b6cb0", label="Payback from today")
    ax[1].plot(startup_df.index, startup_df["Payback from start-up (yr)"], "^--", color="#dd6b20", label="Payback from startup")
    ax[1].set(title="Payback vs Startup Delay", xlabel="Delay (years)", ylabel="Years")
    ax[1].legend()
    plt.tight_layout()
    f5_path = os.path.join(fig_dir, "fig_05_startup_delay.png")
    fig.savefig(f5_path, dpi=160, bbox_inches="tight")
    plt.close(fig)
    generated_figures["startup_delay"] = "figures/fig_05_startup_delay.png"
    
    # Fig 6: Alternatives & Final Comparison
    fig, ax = plt.subplots(1, 2, figsize=(14, 5))
    order_capex = summary_df.sort_values("Capex ($MM)")
    short_lbl = lambda s: s.split(" (")[0]
    ax[0].plot(order_capex["Capex ($MM)"], order_capex["NPV ($MM)"], "o-", color="#1a365d", lw=2)
    for lbl, row in order_capex.iterrows():
        ax[0].annotate(short_lbl(lbl), (row["Capex ($MM)"], row["NPV ($MM)"]), fontsize=8,
                       textcoords="offset points", xytext=(4, 4))
    ax[0].axhline(0, color="black", lw=0.7)
    ax[0].set(title="NPV vs Capital Investment", xlabel="Capex ($MM)", ylabel="NPV ($MM)")
    
    ax[1].bar([short_lbl(l) for l in order_capex.index], order_capex["PIR"], color="#dd6b20")
    ax[1].set(title="Profit / Investment Ratio (PIR = CPI) by Alternative", ylabel="PIR")
    ax[1].tick_params(axis="x", rotation=40)
    plt.tight_layout()
    f6_path = os.path.join(fig_dir, "fig_06_alternatives_comparison.png")
    fig.savefig(f6_path, dpi=160, bbox_inches="tight")
    plt.close(fig)
    generated_figures["alternatives_comparison"] = "figures/fig_06_alternatives_comparison.png"
    
    # Fig 7: Sensitivity Analysis (Spider & Tornado)
    fig, ax = plt.subplots(1, 2, figsize=(14, 5.2))
    for label, (grid, vals) in spider.items():
        ax[0].plot(grid * 100, vals, "o-", label=label, lw=1.8)
    ax[0].axhline(0, color="black", lw=0.7)
    ax[0].axvline(0, color="gray", lw=0.7, ls="--")
    ax[0].set(title=f"Spider Diagram - NPV Sensitivity ({recommendation})", xlabel="Change in Variable (%)", ylabel="NPV ($MM)")
    ax[0].legend(fontsize=8)
    
    base_val = evaluate(build_cash_flow(best_p, cfg=cfg), disc)["NPV"]
    lo_vals = sens_df["NPV @ low"].values - base_val
    hi_vals = sens_df["NPV @ high"].values - base_val
    y_pos = np.arange(len(sens_df))
    ax[1].barh(y_pos, lo_vals, color="#c53030", label="Low Case")
    ax[1].barh(y_pos, hi_vals, color="#2f855a", label="High Case")
    ax[1].axvline(0, color="black", lw=0.8)
    ax[1].set_yticks(y_pos)
    ax[1].set_yticklabels(sens_df.index)
    ax[1].set(title="Tornado Diagram - Impact on NPV ($MM)", xlabel=f"Delta NPV from Base (${base_val:.1f} MM)")
    ax[1].legend(fontsize=8)
    plt.tight_layout()
    f7_path = os.path.join(fig_dir, "fig_07_sensitivity_spider_tornado.png")
    fig.savefig(f7_path, dpi=160, bbox_inches="tight")
    plt.close(fig)
    generated_figures["sensitivity"] = "figures/fig_07_sensitivity_spider_tornado.png"
    
    # 10. Generate Consolidated Excel Workbooks
    results_excel_path = os.path.join(output_dir, "Economic_Indicators_Results.xlsx")
    with pd.ExcelWriter(results_excel_path, engine="openpyxl") as xl:
        base_cf.round(4).to_excel(xl, sheet_name="Base Cash Flow", index=False)
        screening_df.to_excel(xl, sheet_name="Base Indicators", index=False)
        startup_df.round(4).to_excel(xl, sheet_name="Startup Delay")
        summary_df.round(4).to_excel(xl, sheet_name="Alternatives")
        inc_table.round(4).to_excel(xl, sheet_name="Incremental Analysis", index=False)
        comparison.round(4).to_excel(xl, sheet_name="Final Comparison")
        sens_df.round(4).to_excel(xl, sheet_name="Sensitivity")
        pd.DataFrame([{
            "Project": base_inputs.name,
            "Recommendation": recommendation,
            "Discount Rate": disc,
            "Capital Budget": cap_budget,
            "Base NPV ($MM)": base_metrics["NPV"],
            "Recommended NPV ($MM)": summary_df.loc[recommendation, "NPV ($MM)"]
        }]).to_excel(xl, sheet_name="Summary", index=False)
        
    all_cfs_path = os.path.join(output_dir, "All_Cash_Flows.xlsx")
    with pd.ExcelWriter(all_cfs_path, engine="openpyxl") as xl:
        for label, cf_ in alt_cf.items():
            sheet_title = ("Alt " + _slug(label))[:31]
            cf_.round(4).to_excel(xl, sheet_name=sheet_title, index=False)
        for delay, cf_ in startup_cfs.items():
            cf_.round(4).to_excel(xl, sheet_name=f"Delay {delay} yr", index=False)
            
    # Copy original dataset into output folder
    try:
        shutil.copy2(file_path, os.path.join(output_dir, "input_" + os.path.basename(file_path)))
    except Exception as e:
        log(f"Warning: could not copy dataset: {e}")
        
    # Write execution log
    log("All evaluations and workbooks successfully created.")
    log(f"Output files staged at: {output_dir}")
    log_file_path = os.path.join(output_dir, "run_log.txt")
    with open(log_file_path, "w", encoding="utf-8") as lf:
        lf.write(log_stream.getvalue())
        
    # 11. Create Output ZIP Archive
    zip_filename = f"{_slug(stem)}_evaluation_output.zip"
    zip_filepath = os.path.join(output_dir, zip_filename)
    with zipfile.ZipFile(zip_filepath, "w", zipfile.ZIP_DEFLATED) as zf:
        for root, _, files in os.walk(output_dir):
            for file in files:
                if file == zip_filename:
                    continue  # Do not zip the zip itself
                full_path = os.path.join(root, file)
                rel_path = os.path.relpath(full_path, output_dir)
                zf.write(full_path, arcname=rel_path)
                
    log(f"Package ZIP created: {zip_filepath} ({os.path.getsize(zip_filepath):,} bytes)")
    
    # 12. Assemble Results Payload
    return {
        "success": True,
        "run_id": run_id,
        "output_dir": output_dir,
        "zip_filename": zip_filename,
        "zip_filepath": zip_filepath,
        "zip_size_bytes": os.path.getsize(zip_filepath),
        "project": {
            "name": base_inputs.name,
            "life": base_inputs.life,
            "discount_rate": disc,
            "tax_rate": base_inputs.tax,
            "dep_method": base_inputs.dep_method,
            "dep_name": DEP_NAMES[base_inputs.dep_method],
            "dep_years": base_inputs.dep_years,
            "salvage": base_inputs.salvage,
            "oil_price": base_inputs.oil_price,
            "gas_price": base_inputs.gas_price,
            "q1_mbd": base_inputs.q1_mbd,
            "decline": base_inputs.decline,
            "gor": base_inputs.gor,
            "wells": base_inputs.wells,
        },
        "capex_breakdown": capex_dict,
        "base_metrics": {
            "npv": float(base_metrics["NPV"]),
            "irr_text": base_metrics["IRR_text"],
            "irr_roots": [float(r) for r in base_metrics["IRR_roots"]],
            "mirr": float(base_metrics["MIRR"]),
            "payback": float(base_metrics["Payback"]) if base_metrics["Payback"] is not None else None,
            "disc_payback": float(base_metrics["Disc_Payback"]) if base_metrics["Disc_Payback"] is not None else None,
            "pi": float(base_metrics["PI"]),
            "pir": float(base_metrics["PIR"]),
            "pir_undiscounted": float(base_metrics["PIR_undiscounted"]),
            "total_capex": float(base_metrics["Capex"]),
            "last_producing_year": int(base_metrics["Last_Producing_Year"])
        },
        "screening_table": screening_df.to_dict(orient="records"),
        "verification": ver_df.to_dict(orient="records") if ver_df is not None else None,
        "economic_limit": life_df.reset_index().rename(columns={"index": "Case"}).to_dict(orient="records"),
        "startup_delay": startup_df.reset_index().to_dict(orient="records"),
        "alternatives": summary_df.reset_index().to_dict(orient="records"),
        "incremental_analysis": inc_table.to_dict(orient="records"),
        "incremental_winner": inc_winner,
        "best_npv_alternative": best_npv,
        "recommendation": recommendation,
        "final_comparison": comparison.reset_index().to_dict(orient="records"),
        "sensitivity": sens_df.reset_index().to_dict(orient="records"),
        "figures": generated_figures,
        "cash_flow_preview": base_cf.round(3).to_dict(orient="records"),
        "log": log_stream.getvalue()
    }
